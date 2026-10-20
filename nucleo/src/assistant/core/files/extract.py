"""Extracción de texto (o de la imagen reducida) en un subproceso aislado.

El proceso padre llama a `extract(data, kind, settings)`: lanza `python -I -m assistant.core.files.extract` (sin
PYTHONPATH, sin site del usuario y sin el directorio actual en sys.path), le pasa los bytes por stdin y lee un JSON
por stdout. Si tarda más de FILE_EXTRACT_TIMEOUT, se mata. Nada del archivo se ejecuta: pypdf solo lee texto,
openpyxl abre en solo lectura con los valores ya calculados (sin fórmulas ni enlaces externos), python-docx solo
párrafos y tablas, y Pillow verifica la imagen antes de reducirla y la vuelve a codificar sin metadatos (EXIF)."""
import base64
import io
import json
import subprocess
import sys

from assistant.core.files.validate import FileRejected, decode_text

TEXT_PAGE_CHARS = 4000     # un .txt/.md/.csv o un .docx se parte en "páginas" de este tamaño
MAX_IMAGE_PIXELS = 40_000_000


# ------------------------------------------------------------------ proceso padre
def extract(data: bytes, kind: str, settings) -> dict:
    """{"pages": [texto, ...], "chars": n, "truncated": bool, "image": {"b64", "mime", "width", "height"} | None}"""
    cmd = [sys.executable, "-I", "-m", "assistant.core.files.extract", kind, str(settings.file_max_pages),
           str(settings.file_max_chars), str(settings.file_image_max_px)]
    try:
        done = subprocess.run(cmd, input=data, capture_output=True, timeout=settings.file_extract_timeout)
    except subprocess.TimeoutExpired as exc:
        raise FileRejected("extract_timeout") from exc
    if done.returncode != 0:
        try:
            code = json.loads(done.stdout or b"{}").get("error")
        except ValueError:
            code = None
        raise FileRejected(code if code in ("encrypted", "image_too_big", "not_text") else "extract_failed")
    return json.loads(done.stdout)


# ------------------------------------------------------------------ subproceso
def _chunks(text: str) -> list[str]:
    return [text[i:i + TEXT_PAGE_CHARS] for i in range(0, len(text), TEXT_PAGE_CHARS)] or [""]


def _pdf(data: bytes, max_pages: int) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise FileRejected("encrypted")
    return [(page.extract_text() or "") for page in reader.pages[:max_pages]]


def _xlsx(data: bytes, max_pages: int, max_chars: int) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    pages = []
    for ws in wb.worksheets[:max_pages]:
        lines, size = [f"# Sheet: {ws.title}"], 0
        for row in ws.iter_rows(values_only=True):
            line = "\t".join("" if v is None else str(v) for v in row).rstrip()
            if line:
                lines.append(line)
                size += len(line)
            if size > max_chars:
                break
        pages.append("\n".join(lines))
    wb.close()
    return pages


def _docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _image(data: bytes, max_px: int) -> dict:
    from PIL import Image

    Image.MAX_IMAGE_PIXELS = MAX_IMAGE_PIXELS
    try:
        with Image.open(io.BytesIO(data)) as probe:
            probe.verify()
        img = Image.open(io.BytesIO(data))
        img.seek(0)  # GIF animado: solo el primer cuadro
        img.load()
    except Image.DecompressionBombError as exc:
        raise FileRejected("image_too_big") from exc
    img = img.convert("RGBA" if img.mode in ("RGBA", "LA", "P") else "RGB")
    img.thumbnail((max_px, max_px))
    out = io.BytesIO()
    img.save(out, format="PNG")  # PNG nuevo desde los píxeles: sin EXIF ni otros metadatos
    return {"b64": base64.b64encode(out.getvalue()).decode(), "mime": "image/png",
            "width": img.width, "height": img.height}


def run(data: bytes, kind: str, max_pages: int, max_chars: int, max_px: int) -> dict:
    image = None
    if kind == "pdf":
        pages = _pdf(data, max_pages)
    elif kind == "xlsx":
        pages = _xlsx(data, max_pages, max_chars)
    elif kind == "docx":
        pages = _chunks(_docx(data))
    elif kind in ("txt", "md", "csv"):
        pages = _chunks(decode_text(data))
    else:
        pages, image = [], _image(data, max_px)
    truncated = len(pages) > max_pages
    pages, kept, total = pages[:max_pages], [], 0
    for text in pages:
        if total + len(text) > max_chars:
            kept.append(text[: max_chars - total])
            truncated = True
            break
        kept.append(text)
        total += len(text)
    return {"pages": kept, "chars": sum(len(p) for p in kept), "truncated": truncated, "image": image}


def main() -> int:
    kind, max_pages, max_chars, max_px = sys.argv[1], int(sys.argv[2]), int(sys.argv[3]), int(sys.argv[4])
    data = sys.stdin.buffer.read()
    try:
        result = run(data, kind, max_pages, max_chars, max_px)
    except FileRejected as exc:
        _out({"error": exc.code})
        return 2
    except Exception:  # noqa: BLE001 - archivo corrupto o inesperado: el padre solo ve "no se pudo leer"
        _out({"error": "extract_failed"})
        return 3
    _out(result)
    return 0


def _out(obj: dict) -> None:
    # bytes UTF-8: la consola de Windows (cp1252) no puede con todo el texto de un archivo
    sys.stdout.buffer.write(json.dumps(obj, ensure_ascii=False).encode("utf-8"))


if __name__ == "__main__":
    sys.exit(main())
