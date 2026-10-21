"""Archivos que la persona pide descargar: Excel, CSV o el respaldo completo del CRM (JSON). Solo cuando ella lo pide.

El código arma el archivo; el modelo solo elige qué secciones, con qué filtro (los mismos de consultar_datos: `registros.filtrar`)
y en qué formato. Funciones puras: devuelven los bytes; guardarlos y entregarlos es de la ruta (`crm_descargas`).

Formas:
  - `excel`: una hoja por sección, encabezado fijo; números y dinero como números, fechas como fechas (como `exportarExcel.js`).
  - `csv`: UNA sección (CSV no tiene hojas), UTF-8 con BOM para que Excel lea bien los acentos.
  - `respaldo`: el CRM completo en JSON, tal como está guardado (sin `_id` ni `user_id`), para que sus datos nunca queden atrapados.
"""
from __future__ import annotations

import csv
import io
import json
import re
from datetime import date, datetime
from typing import Any

FORMATOS = ("excel", "csv", "respaldo")
MIME = {"excel": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "csv": "text/csv; charset=utf-8",
        "respaldo": "application/json"}
EXTENSION = {"excel": "xlsx", "csv": "csv", "respaldo": "json"}
VERSION_RESPALDO = 1
_FORMATO_NUMERO = {"numero": "#,##0.##", "dinero": "#,##0.00"}
_HOJA_PROHIBIDO = re.compile(r"[\\/?*\[\]:]")
_ARCHIVO_PROHIBIDO = re.compile(r'[\\/:*?"<>|]')


def nombre_archivo(negocio: Any, que: str, formato: str, hoy: date) -> str:
    """«Taller Pepe - Clientes - 2026-10-20.xlsx». Sin caracteres que Windows o macOS no aceptan."""
    partes = [str(x).strip() for x in (negocio or "Mi CRM", que) if str(x or "").strip()]
    base = re.sub(r"\s+", " ", _ARCHIVO_PROHIBIDO.sub(" ", " - ".join(partes))).strip()[:80] or "Mi CRM"
    return f"{base} - {hoy.isoformat()}.{EXTENSION[formato]}"


def _nombre_hoja(nombre: Any, usados: set[str]) -> str:
    """Excel no admite \\ / ? * [ ] : en el nombre de una hoja, ni más de 31 letras, ni dos iguales."""
    base = (re.sub(r"\s+", " ", _HOJA_PROHIBIDO.sub(" ", str(nombre or "Hoja"))).strip() or "Hoja")[:31]
    n, i = base, 2
    while n.lower() in usados:
        n = f"{base[:31 - len(str(i)) - 1]} {i}"
        i += 1
    usados.add(n.lower())
    return n


def _numero(v: Any) -> float | None:
    if isinstance(v, bool) or v is None or str(v).strip() == "":
        return None
    try:
        n = float(v)
    except (TypeError, ValueError):
        return None
    return n if n == n and n not in (float("inf"), float("-inf")) else None


def _fecha(v: Any) -> date | None:
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", str(v or ""))
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3]))
    except ValueError:
        return None


def celda(tipo: str, valor: Any) -> Any:
    """El valor con el tipo que le toca al CAMPO: números y dinero como números, fechas como fechas; lo demás como texto."""
    if valor is None or valor == "":
        return None
    if tipo in _FORMATO_NUMERO and _numero(valor) is not None:
        return _numero(valor)
    if tipo == "fecha" and _fecha(valor):
        return _fecha(valor)
    if isinstance(valor, (list, tuple)):
        return ", ".join(str(x) for x in valor)
    return str(valor)


def tabla(campos: list[dict], registros: list[dict]) -> tuple[list[str], list[list[Any]]]:
    """Encabezados (los nombres que ve la persona) y filas, con los valores ya convertidos por tipo."""
    titulos = [str(c.get("nombre") or c.get("id")) for c in campos]
    filas = [[celda(c.get("tipo", "texto"), (r.get("valores") or {}).get(c["id"])) for c in campos] for r in registros]
    return titulos, filas


def excel(hojas: list[dict]) -> bytes:
    """hojas = [{titulo, campos, registros}] → .xlsx con una hoja por sección."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    libro = Workbook()
    libro.remove(libro.active)
    usados: set[str] = set()
    negrita, fondo = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F6F5B")
    for h in hojas:
        hoja = libro.create_sheet(_nombre_hoja(h["titulo"], usados))
        titulos, filas = tabla(h["campos"], h["registros"])
        hoja.append(titulos)
        for c in hoja[1]:
            c.font, c.fill = negrita, fondo
        for fila in filas:
            hoja.append(fila)
        for i, campo in enumerate(h["campos"], start=1):
            tipo = campo.get("tipo", "texto")
            letra = hoja.cell(row=1, column=i).column_letter
            for (c,) in hoja.iter_rows(min_row=2, min_col=i, max_col=i):
                if tipo in _FORMATO_NUMERO and isinstance(c.value, float):
                    c.number_format = _FORMATO_NUMERO[tipo]
                elif tipo == "fecha" and isinstance(c.value, (date, datetime)):
                    c.number_format = "dd/mm/yyyy"
            largo = max([len(titulos[i - 1])] + [len(str(f[i - 1] or "")) for f in filas[:80]])
            hoja.column_dimensions[letra].width = min(48, max(12, largo + 3))
        hoja.freeze_panes = "A2"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def csv_de(campos: list[dict], registros: list[dict]) -> bytes:
    """Una sección en CSV (UTF-8 con BOM: Excel abre bien los acentos). Las fechas, como AAAA-MM-DD."""
    titulos, filas = tabla(campos, registros)
    salida = io.StringIO()
    w = csv.writer(salida)
    w.writerow(titulos)
    for f in filas:
        w.writerow(["" if v is None else (v.isoformat() if isinstance(v, date) else (int(v) if isinstance(v, float) and v.is_integer() else v)) for v in f])
    return ("﻿" + salida.getvalue()).encode("utf-8")


def respaldo(doc: dict, ahora: datetime) -> bytes:
    """El CRM completo, tal como está guardado (sin `_id` ni `user_id`), con la versión del formato y la fecha."""
    limpio = {k: v for k, v in doc.items() if k not in ("_id", "user_id")}
    contenido = {"formato": "respaldo-crm", "version": VERSION_RESPALDO, "creado": ahora.isoformat(), "crm": limpio}
    return json.dumps(contenido, ensure_ascii=False, indent=1, default=lambda v: v.isoformat() if isinstance(v, (date, datetime)) else str(v)).encode("utf-8")
