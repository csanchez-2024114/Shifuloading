"""Archivos adjuntos contra el agente real: sube txt, csv, docx, xlsx y una imagen, pregunta por su contenido, hace
una segunda pregunta sin adjunto y revisa la puerta de entrada (tipos, tamano, dueno). ~$0.05, ~2 min.
Uso: python scripts/eval_files.py [base]   (por defecto http://localhost:5001; necesita RATE_LIMIT_ENABLED=false)"""
import asyncio
import io
import json
import sys
import uuid
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
sys.argv = [sys.argv[0], str(REPO), sys.argv[1] if len(sys.argv) > 1 else "http://localhost:5001"]
sys.path.insert(0, str(REPO / "scripts"))
sys.path.insert(0, str(REPO))
import eval_agent as ea  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from PIL import Image, ImageDraw  # noqa: E402
from pymongo import MongoClient  # noqa: E402
from dotenv import load_dotenv  # noqa: E402
import os  # noqa: E402

load_dotenv(REPO / ".env")


def make_files() -> dict:
    out = {"notas.txt": "Reunión con el proveedor Harinas del Sur: el saco de 50 kg sube a Q385 desde el 3 de noviembre.".encode("utf-8"),
           "ventas.csv": "producto,unidades,precio\nPastel de chocolate,14,120\nCheesecake,9,150\nGalletas,40,15\n".encode("utf-8")}
    d = Document()
    d.add_paragraph("Contrato de alquiler del local 12, Zona 10. Renta mensual: Q7,500. Vence el 30 de junio de 2027. Arrendador: Inmobiliaria Roble.")
    b = io.BytesIO(); d.save(b); out["contrato.docx"] = b.getvalue()
    wb = Workbook(); ws = wb.active
    for row in (["mes", "gasto"], ["enero", 4200], ["febrero", 3900], ["marzo", 5100]):
        ws.append(row)
    b = io.BytesIO(); wb.save(b); out["gastos.xlsx"] = b.getvalue()
    img = Image.new("RGB", (640, 200), "white")
    ImageDraw.Draw(img).text((20, 80), "FACTURA 00417  TOTAL Q 1,250.00  Cliente: Laura Mendez", fill="black")
    b = io.BytesIO(); img.save(b, "PNG"); out["factura.png"] = b.getvalue()
    return out


# (archivos, pregunta, lo que debe aparecer en la respuesta)
CASES = [
    (["notas.txt"], "a cuánto sube el saco de harina y desde cuándo?", ["385"]),
    (["ventas.csv"], "según el archivo, cuántas unidades de cheesecake se vendieron?", ["9"]),
    (["contrato.docx"], "cuándo vence el contrato y cuánto es la renta?", ["7,500", "2027"]),
    (["gastos.xlsx"], "en qué mes gasté más según la hoja?", ["marzo"]),
    (["factura.png"], "cuál es el total de esta factura?", ["1,250"]),
    (["notas.txt", "contrato.docx"], "resume en dos frases lo que dicen los dos archivos", ["385", "7,500"]),
]


async def turn(a, conv, text, files):
    async with a.http.stream("POST", f"{a.api}/chat/stream", headers={"Authorization": f"Bearer {a.token}"},
                             json={"text": text, "conversation_id": conv, "hold": True, "files": files}, timeout=180) as r:
        async for line in r.aiter_lines():
            if line.startswith("data: "):
                ev = json.loads(line[6:])
                if ev["type"] == "done":
                    return ev["text"]
                if ev["type"] == "error":
                    return f"[ERROR {ev['message']}]"
    return "[sin respuesta]"


async def case(base, blobs, names, ask, must, plan):
    a = ea.Agent(base)
    await a.register()
    h = {"Authorization": f"Bearer {a.token}"}
    if plan:
        import base64
        from bson import ObjectId
        part = a.token.split(".")[1]
        uid = json.loads(base64.urlsafe_b64decode(part + "=" * (-len(part) % 4)))["user_id"]
        n = MongoClient(os.getenv("MONGO_URI"))[os.getenv("MONGO_DB_NAME")].users.update_one({"_id": ObjectId(uid)}, {"$set": {"plan": plan}})
        assert n.matched_count == 1, "no encontre al usuario para subirle el plan"
    ids = []
    for n in names:
        r = await a.http.post(f"{a.api}/files/", headers=h, files={"file": (n, blobs[n])})
        ids.append(r.json().get("file_id") or f"FALLO {r.status_code} {r.text[:80]}")
    conv = uuid.uuid4().hex
    reply = await turn(a, conv, ask, ids)
    again = await turn(a, conv, "y eso, ¿de qué archivo lo sacaste?", [])  # segundo turno sin adjunto: debe recordarlo
    ok = all(m in reply for m in must)
    return f"\n### {names} plan={plan or 'free'} {'OK' if ok else 'FALLA'}\nP: {ask}\nR: {reply[:400]}\nR2: {again[:200]}", ok, (a, h, ids)


async def main():
    base, blobs = sys.argv[2], make_files()
    res = await asyncio.gather(*(case(base, blobs, n, q, m, "pro" if "factura.png" in n else None) for n, q, m in CASES))
    for text, _, _ in res:
        print(text)
    print(f"\nRESPUESTAS CORRECTAS: {sum(ok for _, ok, _ in res)}/{len(res)}")

    # Puerta de entrada: rechazos y dueno
    a, h, ids = res[0][2]
    b = ea.Agent(base); await b.register(); hb = {"Authorization": f"Bearer {b.token}"}
    up = lambda n, d, hh=h: a.http.post(f"{a.api}/files/", headers=hh, files={"file": (n, d)})
    print("exe:", (await up("virus.exe", b"MZ")).status_code, "| pdf falso:", (await up("x.pdf", b"MZ no soy pdf")).status_code,
          "| 11 MB:", (await up("g.txt", b"x" * (11 * 1024 * 1024))).status_code,
          "| sin token:", (await a.http.post(f"{a.api}/files/", files={"file": ("a.txt", b"a")})).status_code,
          "| otro usuario descarga:", (await a.http.get(f"{a.api}/files/{ids[0]}", headers=hb)).status_code,
          "| otro usuario borra:", (await a.http.delete(f"{a.api}/files/{ids[0]}", headers=hb)).status_code,
          "| dueno descarga:", (await a.http.get(f"{a.api}/files/{ids[0]}", headers=h)).status_code)
    ajeno = await turn(b, uuid.uuid4().hex, "a cuánto sube el saco de harina según el archivo adjunto?", ids)
    print("otro usuario adjunta el file_id ajeno ->", "FUGA" if "385" in ajeno else "no lo ve:", ajeno[:160])


asyncio.run(main())
