"""Banco de datos que se conectan (relación + calculados) con la IA real, sobre el chat REAL (/api/v1/crm/chat, CRM_DISENO=1).

Plantilla del gimnasio (Membresías, Pagos.miembro, Clases). Los datos se cargan SIN IA; luego 10 pedidos por el chat, y cada uno se
revisa contra lo que el SERVIDOR calculó (GET /crm), no contra lo que dice el modelo.

    cd backend && CRM_DISENO=1 CRM_TRAZA=1 CRM_TRAZA_DIR=<carpeta> python scripts/crm_dev_server.py --puerto 5213
    BASE=http://localhost:5213 python laboratorio/pruebas/banco_calculos.py

Costo: 10 mensajes (~0,3 M tokens). Los tokens reales salen de la traza.
"""
from __future__ import annotations

import io
import json
import os
import secrets
import sys
import urllib.request

BASE = os.environ.get("BASE", "http://localhost:5213") + "/api/v1"
TOKEN = ""


def pedir(metodo: str, ruta: str, cuerpo: dict | None = None, crudo: bool = False):
    req = urllib.request.Request(BASE + ruta, method=metodo, data=json.dumps(cuerpo).encode() if cuerpo is not None else None,
                                 headers={"Content-Type": "application/json", **({"Authorization": f"Bearer {TOKEN}"} if TOKEN else {})})
    with urllib.request.urlopen(req, timeout=180) as r:
        datos = r.read()
        return datos if crudo else json.loads(datos or b"null")


def lista(crm: dict, lid: str) -> dict:
    return next(l for l in crm["listas"] if l["id"] == lid)


def valor(crm: dict, lid: str, nombre: str, campo_nombre: str):
    l = lista(crm, lid)
    c = next((c for c in l["campos"] if c["nombre"].lower() == campo_nombre.lower()), None)
    if c is None:
        return "SIN CAMPO"
    ct = l["campos"][0]["id"]
    r = next((r for r in l["registros"] if str(r["valores"].get(ct, "")).lower().startswith(nombre.lower())), None)
    return None if r is None else r["valores"].get(c["id"])


def campo(crm: dict, lid: str, pred) -> dict | None:
    return next((c for c in lista(crm, lid)["campos"] if pred(c)), None)


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def esperar(crm, lid, campo_nombre, esperado: dict) -> str | None:
    malos = {n: valor(crm, lid, n, campo_nombre) for n in esperado if _num(valor(crm, lid, n, campo_nombre)) != float(esperado[n])}
    return None if not malos else f"{campo_nombre}: {malos} (esperaba {esperado})"


def calculado_por_nombre(crm, lid, palabras):
    return campo(crm, lid, lambda c: c.get("calculo") and any(p in c["nombre"].lower() for p in palabras))


def caso_enlace(r, crm):
    c = campo(crm, "l_pagos", lambda c: c["id"] == "miembro")
    if c.get("relacion") != "l_membresias":
        return f"Pagos.miembro no quedó enlazado ({c})"
    return None if "pedro" in r["respuesta"].lower() else "no avisó que Pedro Gómez no tiene pareja"


def caso_pagado(r, crm):
    c = calculado_por_nombre(crm, "l_membresias", ["pagad"])
    return "no hay un calculado «Pagado»" if not c else esperar(crm, "l_membresias", c["nombre"], {"Ana": 300, "Luis": 150, "Rosa": 0})


def caso_saldo(r, crm):
    c = calculado_por_nombre(crm, "l_membresias", ["saldo", "debe", "pendiente"])
    return "no hay un calculado de saldo" if not c else esperar(crm, "l_membresias", c["nombre"], {"Ana": 0, "Luis": 150, "Rosa": 250})


def caso_libres(r, crm):
    c = calculado_por_nombre(crm, "l_clases", ["libre", "disponible", "lugar"])
    return "no hay un calculado de lugares libres" if not c else esperar(crm, "l_clases", c["nombre"], {"Yoga": 5, "Spinning": 0})


def caso_total_deuda(r, crm):
    return None if "400" in r["respuesta"].replace(",", "").replace(".00", "") else f"no dijo 400: «{r['respuesta'][:120]}»"


def caso_quien_debe(r, crm):
    return None if "rosa" in r["respuesta"].lower() else f"no dijo Rosa: «{r['respuesta'][:120]}»"


def caso_no_escribe_calculado(r, crm):
    c = calculado_por_nombre(crm, "l_membresias", ["pagad"])
    v = _num(valor(crm, "l_membresias", "Luis", c["nombre"])) if c else None
    if v != 150:
        return f"el pagado de Luis cambió a {v}"
    t = r["respuesta"].lower()
    return f"dice que lo cambió: «{r['respuesta'][:120]}»" if any(x in t for x in ("listo", "ya puse", "cambié", "actualicé")) and "no" not in t[:40] else None


def caso_renombre(r, crm):
    pagos = [x["valores"].get("miembro") for x in lista(crm, "l_pagos")["registros"]]
    if sum(1 for p in pagos if p == "Ana López de Paz") != 2:
        return f"los pagos no siguieron el nombre nuevo: {pagos}"
    c = calculado_por_nombre(crm, "l_membresias", ["pagad"])
    return esperar(crm, "l_membresias", c["nombre"], {"Ana López de Paz": 300}) if c else "sin Pagado"


def caso_dias(r, crm):
    c = campo(crm, "l_membresias", lambda c: (c.get("calculo") or {}).get("op") == "dias_desde")
    return None if c and _num(valor(crm, "l_membresias", "Rosa", c["nombre"])) is not None else "no hay un calculado de días desde el alta"


def caso_excel(r, crm):
    if not r.get("descargas"):
        return "no preparó el Excel"
    from openpyxl import load_workbook

    datos = pedir("GET", f"/crm/descargas/{r['descargas'][0]['id']}", crudo=True)
    h = load_workbook(io.BytesIO(datos)).active
    enc = [c.value for c in h[1]]
    col = next((i for i, x in enumerate(enc) if x and any(p in x.lower() for p in ("saldo", "debe", "pendiente"))), None)
    if col is None:
        return f"el Excel no trae el saldo: {enc}"
    vals = sorted(_num(row[col].value) or 0 for row in h.iter_rows(min_row=2))
    return None if vals == [0, 150, 250] else f"saldos en el Excel: {vals}"


CASOS = [
    ("conecta los pagos con las membresías", caso_enlace),
    ("en membresías agrega un campo Pagado con lo que ha pagado cada miembro", caso_pagado),
    ("y un Saldo que sea el monto menos lo pagado", caso_saldo),
    ("en clases quiero ver los lugares libres: cupo menos inscritos", caso_libres),
    ("¿cuánto me deben en total?", caso_total_deuda),
    ("¿quién me debe más?", caso_quien_debe),
    ("ponle a Luis un pagado de 500", caso_no_escribe_calculado),
    ("cambia el nombre de Ana López a Ana López de Paz", caso_renombre),
    ("agrega a membresías cuántos días lleva cada uno desde su alta", caso_dias),
    ("pásame las membresías en excel", caso_excel),
]


def main() -> int:
    global TOKEN
    r = pedir("POST", "/auth/register", {"name": "Banco", "last_name": "Calculos", "email": f"c{secrets.token_hex(4)}@ejemplo.com", "password": "clave1234x"})
    TOKEN = r["access_token"]
    crm = pedir("POST", "/crm", {"brief": {"nombre": "Gym Fuerza", "rubro": "gimnasio", "equipo": "chico", "canales": ["whatsapp"],
                                           "objetivos": ["cobros"], "que_vendes": "membresías mensuales de gimnasio"}})
    if not {"l_membresias", "l_pagos", "l_clases"} <= {l["id"] for l in crm["listas"]}:
        print("La plantilla del gimnasio cambió: el banco no sirve así.", [l["id"] for l in crm["listas"]])
        return 1
    pedir("POST", "/crm/registros/ejemplos/quitar", {})
    pedir("POST", "/crm/listas/l_membresias/registros/lote", {"filas": [
        {"nombre": "Ana López", "monto": 300, "alta": "2026-09-01"}, {"nombre": "Luis Pérez", "monto": 300, "alta": "2026-10-01"},
        {"nombre": "Rosa Díaz", "monto": 250, "alta": "2026-10-15"}]})
    pedir("POST", "/crm/listas/l_pagos/registros/lote", {"filas": [
        {"miembro": "Ana López", "monto": 200}, {"miembro": "ana lopez", "monto": 100}, {"miembro": "Luis Pérez", "monto": 150},
        {"miembro": "Pedro Gómez", "monto": 50}]})
    pedir("POST", "/crm/listas/l_clases/registros/lote", {"filas": [{"clase": "Yoga", "cupo": 20, "inscritos": 15}, {"clase": "Spinning", "cupo": 10, "inscritos": 10}]})
    ok = 0
    for texto, revisar in CASOS:
        r = pedir("POST", "/crm/chat", {"texto": texto, "contexto": {}})
        crm = pedir("GET", "/crm")
        try:
            error = revisar(r, crm)
        except Exception as e:  # noqa: BLE001
            error = f"{type(e).__name__}: {e}"
        error = error or ("el código le pegó «todavía no está hecho»" if "todavía no está hecho" in r["respuesta"] else None)
        ok += error is None
        print(f"{'✓' if error is None else '✗'} {texto}\n    → {r['respuesta'][:220]}" + (f"\n    ✗ {error}" if error else ""))
    print(f"\n{ok}/{len(CASOS)} casos correctos")
    return 0 if ok == len(CASOS) else 1


if __name__ == "__main__":
    sys.exit(main())
