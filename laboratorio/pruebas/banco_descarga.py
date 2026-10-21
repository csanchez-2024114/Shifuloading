"""Banco de descargas con la IA real, sobre el chat REAL (/api/v1/crm/chat del servidor de desarrollo, con CRM_DISENO=1).

¿Shifu prepara el archivo SOLO cuando se lo piden, con el filtro y el formato correctos, y no dice que lo hizo cuando no?
Cada caso baja el archivo de verdad (GET /crm/descargas/{id}) y revisa su contenido.

    cd backend && CRM_DISENO=1 CRM_TRAZA=1 CRM_TRAZA_DIR=<carpeta> python scripts/crm_dev_server.py --puerto 5211
    BASE=http://localhost:5211 python laboratorio/pruebas/banco_descarga.py

Costo: 10 mensajes (~0,3 M tokens); los datos se cargan sin IA. Los tokens reales salen de la traza (CRM_TRAZA_DIR).
"""
from __future__ import annotations

import csv
import io
import json
import os
import secrets
import sys
import urllib.error
import urllib.request

BASE = os.environ.get("BASE", "http://localhost:5211") + "/api/v1"
TOKEN = ""


def pedir(metodo: str, ruta: str, cuerpo: dict | None = None, crudo: bool = False):
    req = urllib.request.Request(BASE + ruta, method=metodo, data=json.dumps(cuerpo).encode() if cuerpo is not None else None,
                                 headers={"Content-Type": "application/json", **({"Authorization": f"Bearer {TOKEN}"} if TOKEN else {})})
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            datos = r.read()
            return (r.status, datos, dict(r.headers)) if crudo else json.loads(datos or b"null")
    except urllib.error.HTTPError as e:
        if crudo:
            return e.code, e.read(), dict(e.headers)
        raise


def chat(texto: str) -> dict:
    return pedir("POST", "/crm/chat", {"texto": texto, "contexto": {}})


def filas_xlsx(datos: bytes) -> dict[str, int]:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(datos))
    return {h.title: h.max_row - 1 for h in libro.worksheets}


def bajar(r: dict):
    (a,) = r["descargas"]
    estado, datos, cab = pedir("GET", f"/crm/descargas/{a['id']}", crudo=True)
    assert estado == 200, estado
    return a, datos


def sin_archivo(r: dict) -> str | None:
    return None if not r.get("descargas") else f"preparó un archivo que no pidieron: {r['descargas']}"


def _excel(r, hojas_esperadas=None, filas=None):
    if not r.get("descargas"):
        return "no preparó el archivo"
    a, datos = bajar(r)
    if a["formato"] != "excel":
        return f"formato {a['formato']}"
    f = filas_xlsx(datos)
    if hojas_esperadas is not None and (len(f) < hojas_esperadas if hojas_esperadas > 1 else len(f) != 1):
        return f"hojas {f}"
    if filas is not None and sorted(f.values())[-1 if len(f) > 1 else 0] != filas and list(f.values())[0] != filas:
        return f"filas {f} (esperaba {filas})"
    return None


def _csv(r, columnas):
    if not r.get("descargas"):
        return "no preparó el archivo"
    a, datos = bajar(r)
    if a["formato"] != "csv":
        return f"formato {a['formato']}"
    filas = list(csv.reader(io.StringIO(datos.decode("utf-8-sig"))))
    return None if len(filas[0]) == columnas and len(filas) == 6 else f"columnas {filas[0]}, {len(filas) - 1} filas"


def _respaldo(r):
    if not r.get("descargas"):
        return "no preparó el archivo"
    a, datos = bajar(r)
    j = json.loads(datos)
    return None if a["formato"] == "respaldo" and j["formato"] == "respaldo-crm" and j["crm"]["listas"] else f"respaldo raro: {a}"


def _sin_mentir(r, *prohibidas):
    t = r["respuesta"].lower()
    malo = next((p for p in prohibidas if p in t), None)
    return sin_archivo(r) or (f"dice «{malo}»" if malo else None)


CASOS = [
    ("pásame en Excel las reparaciones", lambda r: _excel(r, 1, 5)),
    ("descarga en excel solo las reparaciones que se entregan en octubre", lambda r: _excel(r, 1, 3)),
    ("dame un csv de reparaciones solo con el cliente y el equipo", lambda r: _csv(r, 2)),
    ("hazme un respaldo de todo mi CRM", _respaldo),
    ("exporta todo mi CRM a Excel", lambda r: _excel(r, 2)),
    ("pásame en excel las reparaciones que ya están listas", lambda r: _excel(r, 1, 2)),
    ("¿cuántas reparaciones tengo?", lambda r: sin_archivo(r)),
    ("mándame las reparaciones en PDF", lambda r: _sin_mentir(r, "aquí está tu pdf", "te preparé el pdf", "pdf listo")),
    ("muéstrame las reparaciones", lambda r: sin_archivo(r)),
    ("bájame un csv de todo", lambda r: None if not r.get("descargas") or r["descargas"][0]["formato"] != "csv" else "un CSV con todo (no puede)"),
]

# La plantilla del taller trae «Reparaciones» (l_reparaciones) con equipo, cliente, falla, técnico, estado y entrega (fecha).
FILAS = [{"cliente": c, "equipo": e, "estado": st, "entrega": f} for c, e, st, f in [
    ("Sofía", "laptop HP", "Recibido", "2026-10-02"), ("Marco", "PC", "Listo", "2026-10-10"), ("Ana", "impresora", "Listo", "2026-09-20"),
    ("Luis", "laptop Dell", "Recibido", "2026-10-15"), ("Rosa", "tablet", "Recibido", "2026-09-05")]]


def main() -> int:
    global TOKEN
    correo = f"banco{secrets.token_hex(4)}@ejemplo.com"
    r = pedir("POST", "/auth/register", {"name": "Banco", "last_name": "Descargas", "email": correo, "password": "clave1234x"})
    TOKEN = r["access_token"]
    pedir("POST", "/crm", {"brief": {"nombre": "Taller Pepe", "rubro": "reparación de computadoras", "equipo": "chico", "canales": ["whatsapp"],
                                     "objetivos": ["seguimientos"], "que_vendes": "reparación de computadoras, laptops e impresoras"}})
    pedir("POST", "/crm/registros/ejemplos/quitar", {})  # sin los ejemplos de la plantilla: solo los 5 de abajo
    r = pedir("POST", "/crm/listas/l_reparaciones/registros/lote", {"filas": FILAS})
    regs = next((len(s.get("registros") or []) for s in r.get("listas") or [] if s["id"] == "l_reparaciones"), 0)
    print(f"Preparación (sin IA): {regs} reparaciones.")
    if regs != 5:
        print("No quedaron las 5 reparaciones (¿cambió la plantilla del taller?): el banco no sirve así.")
        return 1
    ok = 0
    for texto, revisar in CASOS:
        r = chat(texto)
        try:
            error = revisar(r) or ("el código le pegó «todavía no está hecho» a algo hecho" if "todavía no está hecho" in r["respuesta"] else None)
        except Exception as e:  # noqa: BLE001
            error = f"{type(e).__name__}: {e}"
        ok += error is None
        print(f"{'✓' if error is None else '✗'} {texto}\n    → {r['respuesta'][:200]}" + (f"\n    ✗ {error}" if error else ""))
    print(f"\n{ok}/{len(CASOS)} casos correctos")
    return 0 if ok == len(CASOS) else 1


if __name__ == "__main__":
    sys.exit(main())
